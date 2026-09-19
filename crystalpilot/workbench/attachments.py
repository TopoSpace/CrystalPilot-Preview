"""Composer attachments: turn uploaded files into agent-turn input.

Images become true multimodal input (codex ``LocalImageInput`` - the model
sees the pixels). Documents (PDF/Word/plain text) are referenced by project
path AND, when extractable, inlined as a bounded text excerpt so the model
can use them immediately without a shell round-trip. Anything else is
referenced by path only.

Security: rel paths from the client are confined to the project directory
(no `..`, no absolute paths, must resolve inside the project root).
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

# per-file excerpt bound: enough for a few paper pages / an SOP, small
# enough not to blow the turn context when several files are attached
MAX_EXCERPT_CHARS = 12_000
MAX_PDF_PAGES = 40

_TEXT_SUFFIXES = {".txt", ".md", ".csv", ".log", ".json", ".yaml", ".yml",
                  ".ins", ".res", ".cif", ".fcf", ".hkl", ".pcf", ".lst",
                  ".out", ".prp", ".abs"}
_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}


def _cap(text: str, note_path: str) -> str:
    text = text.strip()
    if len(text) > MAX_EXCERPT_CHARS:
        return (text[:MAX_EXCERPT_CHARS]
                + f"\n…（已截断；完整文件在项目内 {note_path}，可直接读取）")
    return text


def extract_text(path: Path) -> str | None:
    """Best-effort text extraction; None when the format has no extractor
    or extraction fails (caller falls back to a path reference)."""
    suffix = path.suffix.lower()
    rel_note = path.name
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            pages = reader.pages[:MAX_PDF_PAGES]
            parts = []
            for pg in pages:
                t = pg.extract_text() or ""
                if t.strip():
                    parts.append(t)
            if not parts:
                return None
            body = "\n\n".join(parts)
            if len(reader.pages) > MAX_PDF_PAGES:
                body += f"\n…（仅抽取前 {MAX_PDF_PAGES} 页，共 {len(reader.pages)} 页）"
            return _cap(body, rel_note)
        if suffix == ".docx":
            import docx
            doc = docx.Document(str(path))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
            if not parts:
                return None
            return _cap("\n".join(parts), rel_note)
        if suffix in _TEXT_SUFFIXES:
            raw = path.read_bytes()[: MAX_EXCERPT_CHARS * 4]
            return _cap(raw.decode("utf-8", errors="replace"), rel_note)
    except Exception:  # noqa: BLE001 - extraction is best-effort by design
        return None
    return None


def resolve_attachments(
    project_root: Path,
    attachments: list[dict[str, Any]] | None,
) -> tuple[list[str], list[str], list[str]]:
    """Validate client attachment refs against the project directory.

    Returns (image_paths, context_lines, names):
    - image_paths: absolute paths for codex LocalImageInput;
    - context_lines: text appended to the user message (path references and
      document excerpts) - also what makes the transcript self-describing;
    - names: display names for event logs.
    """
    images: list[str] = []
    lines: list[str] = []
    names: list[str] = []
    root = project_root.resolve()
    for a in attachments or []:
        rel = str(a.get("rel", "")).replace("\\", "/").strip()
        if rel == "" or rel.startswith("/") or ".." in rel.split("/"):
            lines.append(f"[附件无效：{rel or '(空路径)'}]")
            continue
        p = (root / rel).resolve()
        if root not in p.parents or not p.is_file():
            lines.append(f"[附件缺失：{rel}]")
            continue
        name = str(a.get("name") or p.name)
        names.append(name)
        suffix = p.suffix.lower()
        if suffix in _IMAGE_SUFFIXES:
            images.append(str(p))
            lines.append(f"[图片附件] {rel}（{name}，已作为图像输入随本条消息发送）")
            continue
        excerpt = extract_text(p)
        if excerpt is not None:
            kind = {".pdf": "PDF", ".docx": "Word"}.get(suffix, "文本")
            lines.append(
                f"[附件] {rel}（{kind}），以下为自动抽取的文本内容，"
                f"原文件在项目目录内可直接读取：\n<<<\n{excerpt}\n>>>")
        else:
            lines.append(f"[附件] {rel}（二进制或无法抽取文本；"
                         f"文件已存放在项目目录内，可用工具读取）")
    return images, lines, names


def compose_message(message: str,
                    context_lines: list[str]) -> str:
    if not context_lines:
        return message
    tail = "\n\n".join(context_lines)
    return f"{message}\n\n{tail}" if message.strip() else tail
