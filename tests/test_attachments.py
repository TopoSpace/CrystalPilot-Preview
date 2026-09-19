"""Composer attachments: path confinement, extraction, message assembly.
Pure unit tests - no live agent, no codex."""
from pathlib import Path

from crystalpilot.workbench.attachments import (compose_message, extract_text,
                                                resolve_attachments)


def _mk(tmp_path: Path, rel: str, data: bytes = b"x") -> Path:
    p = tmp_path / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_bytes(data)
    return p


class TestResolve:
    def test_image_goes_to_multimodal_input(self, tmp_path):
        _mk(tmp_path, "uploads/shot.png", b"\x89PNG\r\n\x1a\nrest")
        images, lines, names = resolve_attachments(
            tmp_path, [{"rel": "uploads/shot.png", "kind": "image"}])
        assert len(images) == 1
        assert images[0].endswith("shot.png")
        assert names == ["shot.png"]
        assert any("图片附件" in ln for ln in lines)

    def test_traversal_rejected(self, tmp_path):
        _mk(tmp_path, "uploads/a.txt")
        for evil in ("../secrets.txt", "uploads/../../x", "/abs/path.png",
                     "C:/Windows/x.png"):
            images, lines, _ = resolve_attachments(
                tmp_path, [{"rel": evil, "kind": "other"}])
            assert images == []
            assert len(lines) == 1
            assert "附件无效" in lines[0] or "附件缺失" in lines[0]

    def test_missing_file_noted_not_fatal(self, tmp_path):
        images, lines, names = resolve_attachments(
            tmp_path, [{"rel": "uploads/nope.pdf", "kind": "other"}])
        assert images == [] and names == []
        assert "附件缺失" in lines[0]

    def test_text_file_inlined(self, tmp_path):
        _mk(tmp_path, "uploads/notes.txt", "晶体在 100K 收集".encode())
        images, lines, _ = resolve_attachments(
            tmp_path, [{"rel": "uploads/notes.txt", "kind": "other"}])
        assert images == []
        joined = "\n".join(lines)
        assert "晶体在 100K 收集" in joined
        assert "uploads/notes.txt" in joined

    def test_binary_referenced_by_path_only(self, tmp_path):
        _mk(tmp_path, "uploads/data.bin", b"\x00\x01\x02")
        _, lines, _ = resolve_attachments(
            tmp_path, [{"rel": "uploads/data.bin", "kind": "other"}])
        assert "无法抽取文本" in lines[0]

    def test_display_name_override(self, tmp_path):
        _mk(tmp_path, "uploads/img_1.png", b"\x89PNG\r\n\x1a\n")
        _, _, names = resolve_attachments(
            tmp_path, [{"rel": "uploads/img_1.png", "kind": "image",
                        "name": "反应方案.png"}])
        assert names == ["反应方案.png"]


class TestExtract:
    def test_txt_capped(self, tmp_path):
        big = _mk(tmp_path, "big.md", b"A" * 100_000)
        out = extract_text(big)
        assert out is not None
        assert len(out) < 13_000
        assert "截断" in out

    def test_docx_paragraphs_and_tables(self, tmp_path):
        import docx
        doc = docx.Document()
        doc.add_paragraph("第一段：样品来自反应釜 3。")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "温度"
        table.rows[0].cells[1].text = "100K"
        p = tmp_path / "note.docx"
        doc.save(str(p))
        out = extract_text(p)
        assert out is not None
        assert "反应釜 3" in out
        assert "温度 | 100K" in out

    def test_pdf_text(self, tmp_path):
        from pypdf import PdfWriter
        w = PdfWriter()
        w.add_blank_page(width=200, height=200)
        p = tmp_path / "blank.pdf"
        with open(p, "wb") as fh:
            w.write(fh)
        # blank page -> no extractable text -> None (falls back to path ref)
        assert extract_text(p) is None

    def test_unknown_suffix_none(self, tmp_path):
        p = _mk(tmp_path, "frame.sfrm", b"\x00" * 10)
        assert extract_text(p) is None


class TestCompose:
    def test_plain_passthrough(self):
        assert compose_message("hi", []) == "hi"

    def test_appends_lines(self):
        out = compose_message("解析这个", ["[图片附件] uploads/a.png（a.png…）"])
        assert out.startswith("解析这个\n\n[图片附件]")

    def test_attachment_only_message(self):
        out = compose_message("", ["[附件] uploads/a.txt"])
        assert out == "[附件] uploads/a.txt"
